from __future__ import annotations

import subprocess
import tempfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pytesseract
from docx import Document as DocxDocument
from pypdf import PdfReader
from unstructured.partition.auto import partition


SUPPORTED_MIME_TO_SUFFIX = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/msword": ".doc",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "text/html": ".html",
    "application/xhtml+xml": ".html",
    "application/epub+zip": ".epub",
}

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".doc", ".txt", ".md", ".xlsx", ".pptx", ".html", ".epub"}


@dataclass(slots=True)
class TextElement:
    text: str
    page_num: int | None
    element_type: str  # title | paragraph | table | list


class DocumentProcessor:
    def parse(self, raw: bytes, mime_type: str | None) -> list[TextElement]:
        suffix = self._suffix_from_mime(mime_type)
        if suffix not in SUPPORTED_SUFFIXES:
            raise ValueError(f"Unsupported file type: {mime_type}")

        # Primary parser for all supported formats.
        elements = self._parse_with_unstructured(raw, suffix)

        # PDF fallback parser.
        if not elements and suffix == ".pdf":
            elements = self._parse_pdf_fallback(raw)

        # DOCX fallback parser.
        if not elements and suffix == ".docx":
            elements = self._parse_docx_fallback(raw)

        # DOC fallback: convert to DOCX via LibreOffice, then parse with python-docx.
        if not elements and suffix == ".doc":
            elements = self._parse_doc_fallback(raw)

        # Scanned PDF OCR fallback when average extracted chars/page is too low.
        if suffix == ".pdf" and self._needs_pdf_ocr(raw, elements):
            ocr_elements = self._parse_pdf_with_ocr(raw)
            if ocr_elements:
                return ocr_elements

        return elements

    @staticmethod
    def _suffix_from_mime(mime_type: str | None) -> str:
        if not mime_type:
            raise ValueError("mime_type is required")
        mime = mime_type.split(";", 1)[0].strip().lower()
        suffix = SUPPORTED_MIME_TO_SUFFIX.get(mime)
        if not suffix:
            raise ValueError(f"Unsupported file type: {mime_type}")
        return suffix

    def _parse_with_unstructured(self, raw: bytes, suffix: str) -> list[TextElement]:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(raw)
            tmp_path = Path(tmp.name)

        try:
            parsed = partition(filename=str(tmp_path))
            return self._elements_from_unstructured(parsed)
        except Exception:
            return []
        finally:
            tmp_path.unlink(missing_ok=True)

    @staticmethod
    def _elements_from_unstructured(parsed_elements: list[object]) -> list[TextElement]:
        out: list[TextElement] = []
        for element in parsed_elements:
            text = (getattr(element, "text", "") or "").strip()
            if not text:
                continue

            metadata = getattr(element, "metadata", None)
            page_num = getattr(metadata, "page_number", None) if metadata is not None else None
            category = getattr(element, "category", "") or element.__class__.__name__

            out.append(
                TextElement(
                    text=text,
                    page_num=page_num,
                    element_type=DocumentProcessor._map_element_type(category),
                )
            )
        return out

    @staticmethod
    def _map_element_type(category: str) -> str:
        c = category.lower()
        if "title" in c:
            return "title"
        if "table" in c:
            return "table"
        if "list" in c:
            return "list"
        return "paragraph"

    @staticmethod
    def _parse_pdf_fallback(raw: bytes) -> list[TextElement]:
        reader = PdfReader(BytesIO(raw))
        elements: list[TextElement] = []
        for idx, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                elements.append(TextElement(text=text, page_num=idx, element_type="paragraph"))
        return elements

    @staticmethod
    def _parse_docx_fallback(raw: bytes) -> list[TextElement]:
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(raw)
            docx_path = Path(tmp.name)

        try:
            doc = DocxDocument(str(docx_path))
            elements: list[TextElement] = []
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = (paragraph.style.name if paragraph.style else "") or ""
                lowered = style_name.lower()
                if "heading" in lowered or "title" in lowered:
                    element_type = "title"
                elif "list" in lowered:
                    element_type = "list"
                else:
                    element_type = "paragraph"
                elements.append(TextElement(text=text, page_num=None, element_type=element_type))

            for table in doc.tables:
                rows: list[str] = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    row_text = " | ".join(cell for cell in cells if cell)
                    if row_text:
                        rows.append(row_text)
                if rows:
                    elements.append(TextElement(text="\n".join(rows), page_num=None, element_type="table"))
            return elements
        finally:
            docx_path.unlink(missing_ok=True)

    def _parse_doc_fallback(self, raw: bytes) -> list[TextElement]:
        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            doc_path = tmp_path / "input.doc"
            docx_path = tmp_path / "input.docx"
            doc_path.write_bytes(raw)

            command = [
                "soffice",
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                str(tmp_path),
                str(doc_path),
            ]
            try:
                subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            except FileNotFoundError as exc:
                raise RuntimeError("LibreOffice (soffice) is required to parse DOC files") from exc
            except subprocess.CalledProcessError as exc:
                stderr = exc.stderr.decode("utf-8", errors="ignore")
                raise RuntimeError(f"DOC to DOCX conversion failed: {stderr}") from exc

            if not docx_path.exists():
                raise RuntimeError("DOC to DOCX conversion failed: output file not found")

            return self._parse_docx_fallback(docx_path.read_bytes())

    @staticmethod
    def _needs_pdf_ocr(raw: bytes, elements: list[TextElement]) -> bool:
        try:
            reader = PdfReader(BytesIO(raw))
            page_count = len(reader.pages)
        except Exception:
            page_count = 0

        if page_count <= 0:
            return False

        total_chars = sum(len(element.text.strip()) for element in elements)
        avg_chars_per_page = total_chars / page_count
        return avg_chars_per_page < 50

    @staticmethod
    def _parse_pdf_with_ocr(raw: bytes) -> list[TextElement]:
        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp_dir:
            tmp_path = Path(tmp_dir)
            pdf_path = tmp_path / "input.pdf"
            pdf_path.write_bytes(raw)

            output_prefix = tmp_path / "page"
            command = ["pdftoppm", "-png", str(pdf_path), str(output_prefix)]
            try:
                subprocess.run(command, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            except (FileNotFoundError, subprocess.CalledProcessError):
                return []

            image_paths = sorted(tmp_path.glob("page-*.png"))
            elements: list[TextElement] = []
            for idx, image_path in enumerate(image_paths, start=1):
                with Image.open(image_path) as image:
                    text = pytesseract.image_to_string(image).strip()
                if text:
                    elements.append(TextElement(text=text, page_num=idx, element_type="paragraph"))

            return elements


document_processor = DocumentProcessor()
